from docx import Document
from .base import ExtractionResult, IngestionError

def extract_docx(path: str) -> ExtractionResult:
    try:
        doc = Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables = []
        for table in doc.tables:
            tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
        text_parts = paragraphs[:]
        for table in tables:
            text_parts.append("TABLE:\n" + "\n".join(" | ".join(row) for row in table))
        return ExtractionResult("\n".join(text_parts), {"paragraph_count": len(paragraphs), "table_count": len(tables), "tables": tables})
    except Exception as exc:
        raise IngestionError(f"DOCX extraction failed: {exc}") from exc
